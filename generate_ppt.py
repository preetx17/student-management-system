import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    img_dir = os.path.join(os.getcwd(), "diagrams")

    # High-Contrast Minimalist Theme
    DARK_BG = RGBColor(7, 10, 18)           
    CARD_BG = RGBColor(17, 24, 39)          
    PRIMARY = RGBColor(99, 102, 241)        
    SECONDARY = RGBColor(16, 185, 129)      
    ACCENT_CYAN = RGBColor(0, 242, 254)     
    ACCENT_ROSE = RGBColor(244, 63, 94)     
    WHITE = RGBColor(255, 255, 255)
    LIGHT_TEXT = RGBColor(226, 232, 240)    
    MUTED = RGBColor(148, 163, 184)         

    def set_bg(slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = DARK_BG

    def add_hero(slide, title, subtitle, badge_text="FINAL PRESENTATION"):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(1.5), Inches(3.1), Inches(0.5))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PRIMARY
        badge.line.fill.background()
        p_b = badge.text_frame.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(14)
        p_b.font.bold = True
        p_b.font.color.rgb = WHITE
        p_b.alignment = PP_ALIGN.CENTER

        tbox = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.33), Inches(2.0))
        tf = tbox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(64)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(28)
        p2.font.color.rgb = ACCENT_CYAN
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(20)

    def add_split(slide, badge_text, title, subtitle, list_title, items, is_right=False):
        # Left Side (or right)
        txt_left = Inches(7.2) if is_right else Inches(0.8)
        
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, txt_left, Inches(1.8), Inches(2.5), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = DARK_BG
        badge.line.color.rgb = PRIMARY
        badge.line.width = Pt(1.5)
        p_b = badge.text_frame.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = PRIMARY
        p_b.alignment = PP_ALIGN.CENTER

        t_box = slide.shapes.add_textbox(txt_left, Inches(2.5), Inches(5.3), Inches(2.0))
        tf = t_box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(44)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(22)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(10)

        # Card Side
        card_left = Inches(0.8) if is_right else Inches(6.8)
        add_card(slide, card_left, Inches(1.5), Inches(5.7), Inches(4.5), list_title, items, font_size=18)

    def add_card(slide, left, top, width, height, title="", items=[], bg_color=CARD_BG, border_color=ACCENT_CYAN, font_size=18):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.4), top + Inches(0.4), width - Inches(0.8), height - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        if title:
            p_t = tf.paragraphs[0]
            p_t.text = title
            p_t.font.size = Pt(font_size + 6)
            p_t.font.bold = True
            p_t.font.color.rgb = WHITE
            p_t.space_after = Pt(20)

        for i, item in enumerate(items):
            p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
            p.text = f"→  {item}"
            p.font.size = Pt(font_size)
            p.font.color.rgb = LIGHT_TEXT
            p.space_after = Pt(16)

    def add_full_image_slide(slide, title, category, image_filename):
        # Small header
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(14)
        p_c.font.bold = True
        p_c.font.color.rgb = PRIMARY

        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf = t_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = WHITE

        # Image
        img_path = os.path.join(img_dir, image_filename)
        if os.path.exists(img_path):
            # Center large image
            slide.shapes.add_picture(img_path, Inches(0.8), Inches(1.8), Inches(11.73), Inches(5.2))
        else:
            # Placeholder box if image is missing
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.73), Inches(5.2))
            shape.fill.solid()
            shape.fill.fore_color.rgb = MUTED
            tf = shape.text_frame
            tf.text = f"Please insert screenshot here:\n{image_filename}"
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            tf.paragraphs[0].font.size = Pt(30)


    # ================== SLIDE 1: HERO ==================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)
    add_hero(s1, "Student Management System", "A Secure, Scalable, and Modern Platform")

    # ================== SLIDE 2: THE PROBLEM ==================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_split(s2, "The Old Way", "Why Upgrade?", "Legacy systems were holding operations back with fragmented data and slow UI.", "Key Limitations", [
        "Vulnerable Data: Excel and paper logs lead to data loss and corruption.",
        "Device Locked: No cloud access meant being stuck to one office PC.",
        "Zero Role Security: Everyone had unrestricted access to all sensitive data.",
        "Slow Manual Workflow: Searching and updating records was extremely tedious."
    ])

    # ================== SLIDE 3: THE SOLUTION ==================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_split(s3, "The New Way", "A Modern Solution", "We built an automated, secure, and beautiful centralized platform.", "Major Upgrades", [
        "Real-time Cloud Sync: Access records securely from any device.",
        "Strict Role Hierarchies: Owner vs Teacher privileges ensure data safety.",
        "Bulk Automation: 1-Click CSV Import and Export for hundreds of records.",
        "Elegant SPA Design: Glassmorphism UI that feels premium and responsive."
    ], is_right=True)

    # ================== SLIDE 4: ARCHITECTURE ==================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    cat_box = s4.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12.33), Inches(0.4))
    p_c = cat_box.text_frame.paragraphs[0]
    p_c.text = "TECH STACK"
    p_c.font.size = Pt(14)
    p_c.font.bold = True
    p_c.font.color.rgb = PRIMARY
    p_c.alignment = PP_ALIGN.CENTER

    t_box = s4.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12.33), Inches(0.8))
    p = t_box.text_frame.paragraphs[0]
    p.text = "Under The Hood"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    add_card(s4, Inches(0.5), Inches(2.8), Inches(2.8), Inches(4.2), "Frontend", ["HTML5 & CSS3 Modules", "Vanilla JS for lightning speed", "Premium Glassmorphism"], font_size=16)
    add_card(s4, Inches(3.6), Inches(2.8), Inches(2.8), Inches(4.2), "Backend", ["Node.js Environment", "Express 5 API Router", "Stateless Session Management"], font_size=16)
    add_card(s4, Inches(6.7), Inches(2.8), Inches(2.8), Inches(4.2), "Database", ["Serverless MySQL Engine", "Bcrypt Hash Security", "Normalized Tables"], font_size=16)
    add_card(s4, Inches(9.8), Inches(2.8), Inches(2.8), Inches(4.2), "Deployment", ["Automated CI/CD Pipeline", "Zero-downtime Rollouts", "Environment Secrets"], font_size=16)

    # ================== SLIDE 5: ERD ==================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_full_image_slide(s5, "Database Entities & Relationships", "DATABASE", "er_diagram.png")

    # ================== SLIDE 6: DFD L0 ==================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_full_image_slide(s6, "System Boundaries (Context)", "DFD 0", "dfd_level0_diagram.png")

    # ================== SLIDE 7: DFD L1 ==================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_full_image_slide(s7, "Core Processes & Decomposed Modules", "DFD 1", "dfd_level1_diagram.png")

    # ================== SLIDE 8: OLD UI (Middle of Presentation) ==================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_full_image_slide(s8, "The Previous Dashboard UI", "BEFORE EVOLUTION", "before_dashboard.png")
    
    # ================== SLIDE 9: OLD UI 2 (Middle of Presentation) ==================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_full_image_slide(s9, "Previous Form Inputs & Tables", "BEFORE EVOLUTION", "before_forms.png")

    # ================== SLIDE 10: NEW UI HIGHLIGHTS ==================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    
    cat_box = s10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.4))
    p_c = cat_box.text_frame.paragraphs[0]
    p_c.text = "THE TRANSFORMATION"
    p_c.font.size = Pt(14)
    p_c.font.bold = True
    p_c.font.color.rgb = PRIMARY
    
    t_box = s10.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.8))
    p = t_box.text_frame.paragraphs[0]
    p.text = "What Did We Actually Change?"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE

    add_card(s10, Inches(0.8), Inches(2.8), Inches(3.6), Inches(4.0), "Single Page App", ["Complete rewrite to remove page reloads.", "Provides an 'App-like' instant feel.", "Dynamic DOM manipulation."], font_size=16)
    add_card(s10, Inches(4.8), Inches(2.8), Inches(3.6), Inches(4.0), "Integrated Views", ["Created unified Teacher and Student rosters.", "Advanced course filtering added.", "Clean, modern data tables."], font_size=16)
    add_card(s10, Inches(8.8), Inches(2.8), Inches(3.6), Inches(4.0), "Quick Modals", ["Add/Edit records without leaving the page.", "Real-time visual toast notifications.", "Responsive on all screen sizes."], font_size=16)

    # ================== SLIDE 11: SECURITY ==================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11)
    add_split(s11, "Security", "Rock Solid Defense", "Enterprise-grade protection was built into every layer.", "Defense Mechanisms", [
        "Cryptographic Hashing: Bcrypt prevents reverse-engineering of passwords.",
        "Network Guards: Session tokens and IP rate limiting block brute-force attacks.",
        "Data Sanitization: Malicious SQL payloads are neutralized instantly via prepared statements."
    ])

    # ================== SLIDE 12: WORKFLOWS ==================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12)
    add_full_image_slide(s12, "Process Logic & Validations", "WORKFLOWS", "process_flow_diagram.png")

    # ================== SLIDE 13: USE CASE ==================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13)
    add_full_image_slide(s13, "Actor Privileges & Access", "USE CASE", "usecase_diagram.png")

    # ================== SLIDE 14: SUMMARY ==================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14)
    add_hero(s14, "Built for the Future", "A massive leap to a secure, automated, scalable platform.", "SUMMARY")

    # ================== SLIDE 15: Q&A ==================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15)
    add_hero(s15, "Thank You", "Open for Questions & Discussion", "Q & A")

    # Save
    try:
        prs.save("Student_Management_System_Presentation_v4.pptx")
        print("Successfully saved presentation v4.")
    except Exception as e:
        print(f"Error saving: {e}")

if __name__ == "__main__":
    build_presentation()
